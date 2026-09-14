"""atualiza Apresentacao_T1_OpenMP.pptx com figuras e textos mais explicativos."""
from __future__ import annotations

import shutil
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "Apresentacao_T1_OpenMP.pptx"
BACKUP = ROOT / "apresentacao" / "Apresentacao_T1_OpenMP_antes.pptx"
FIG = ROOT / "apresentacao" / "figuras"

NAVY = RGBColor(0x0B, 0x21, 0x38)
TEAL = RGBColor(0x02, 0xC3, 0x9A)
BLUE = RGBColor(0x1C, 0x72, 0x93)
INK = RGBColor(0x1E, 0x29, 0x33)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xD9, 0xE2, 0xEC)
LINE = RGBColor(0xE3, 0xEA, 0xF0)


def delete_shape(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def set_run(run, text, size, bold, name, color):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run.font.color.rgb = color


def fill_tf(tf, lines, default_size=15, default_name="Calibri", default_color=INK, default_bold=False):
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, str):
            item = {"text": item}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = item.get("align")
        p.space_after = Pt(item.get("after", 6))
        run = p.add_run()
        set_run(
            run,
            item["text"],
            item.get("size", default_size),
            item.get("bold", default_bold),
            item.get("name", default_name),
            item.get("color", default_color),
        )


def add_box(slide, l, t, w, h, lines, **kwargs):
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    fill_tf(box.text_frame, lines, **kwargs)
    return box


def add_rect(slide, l, t, w, h, color, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def add_header(slide, number, kicker, title):
    oval = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Emu(822960), Emu(685800), Emu(566928), Emu(566928)
    )
    oval.fill.solid()
    oval.fill.fore_color.rgb = TEAL
    oval.line.fill.background()
    num = slide.shapes.add_textbox(Emu(822960), Emu(685800), Emu(566928), Emu(566928))
    tf = num.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    set_run(run, str(number), 24, True, "Cambria", NAVY)
    add_box(
        slide, 1600200, 749808, 9800000, 320040,
        [{"text": kicker, "size": 12, "bold": True, "name": "Calibri", "color": BLUE, "after": 0}],
    )
    add_box(
        slide, 1600200, 1097280, 10241280, 820000,
        [{"text": title, "size": 26, "bold": True, "name": "Cambria", "color": INK, "after": 0}],
    )


def set_badge_number(slide, n):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and len(sh.text_frame.text.strip()) <= 2:
            fill_tf(
                sh.text_frame,
                [{"text": str(n), "size": 24, "bold": True, "name": "Cambria",
                  "color": NAVY, "align": PP_ALIGN.CENTER, "after": 0}],
            )
            return


def insert_slide(prs, index, layout):
    slide = prs.slides.add_slide(layout)
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst[-1]
    sldIdLst.remove(sldId)
    sldIdLst.insert(index, sldId)
    return slide


def wipe_body(slide, keep=4):
    for sh in list(slide.shapes)[keep:]:
        delete_shape(sh)


def add_pic(slide, path, l, t, w, h):
    return slide.shapes.add_picture(str(path), Emu(l), Emu(t), Emu(w), Emu(h))


def blank_slide(prs, index, number, kicker, title):
    slide = insert_slide(prs, index, prs.slide_layouts[0])
    for sh in list(slide.shapes):
        delete_shape(sh)
    add_header(slide, number, kicker, title)
    return slide


def read_ppm(path):
    tokens = Path(path).read_text().split()
    w, h = int(tokens[1]), int(tokens[2])
    rgb = np.array([int(x) for x in tokens[4:]], dtype=np.uint8)
    return rgb.reshape(h, w, 3)


def ensure_render_pngs():
    mapping = {
        "render_spp1.ppm": "render_spp1.png",
        "render_spp4.ppm": "render_spp4.png",
        "render_spp32.ppm": "render_spp32.png",
    }
    for src, dst in mapping.items():
        out = FIG / dst
        if not out.exists():
            Image.fromarray(read_ppm(FIG / src)).save(out)


def update_aplicacao(slide):
    wipe_body(slide, keep=4)
    add_box(
        slide, 822960, 1980000, 5200000, 4500000,
        [
            {"text": "path tracer mínimo do kevin beason (2008), reescrito em c11. não é um motor de jogos: é um integrador de luz.",
             "size": 14, "name": "Calibri", "color": INK, "after": 8},
            {"text": "para cada pixel, o programa dispara raios da câmera, deixa ricochetear na cena e soma a cor da luz que o caminho coletou. um único caminho é ruidoso; a média de muitos caminhos vira a imagem.",
             "size": 14, "name": "Calibri", "color": INK, "after": 8},
            {"text": "a cena é a cornell box: parede vermelha, parede azul, esfera de espelho, esfera de vidro e luz no teto. as “paredes” são esferas gigantes.",
             "size": 14, "name": "Calibri", "color": INK, "after": 8},
            {"text": "entrada: -w largura, -h altura, -s spp (amostras por subpixel). custo proporcional a w · h · spp. referência: 800×600, spp=32 → T1 ≈ 54 s.",
             "size": 14, "name": "Calibri", "color": INK, "after": 0},
        ],
    )
    add_pic(slide, FIG / "17_render_referencia_800x600_spp32.png",
            6200000, 1980000, 5400000, 3200000)
    add_rect(slide, 6200000, 5300000, 5400000, 1200000, NAVY)
    add_box(
        slide, 6400000, 5420000, 5000000, 1000000,
        [
            {"text": "por que este programa    T1 ≈ 54 s    1 laço por linha    não trivial (espelho, vidro, monte carlo)",
             "size": 13, "name": "Calibri", "color": WHITE, "after": 0},
        ],
    )


def make_como_funciona(prs):
    slide = blank_slide(prs, 2, 2, "A APLICAÇÃO", "Como um pixel ganha cor")
    add_pic(slide, FIG / "04_path_tracing_rotulado.png",
            700000, 1980000, 6400000, 4400000)
    add_box(
        slide, 7300000, 1980000, 4400000, 4500000,
        [
            {"text": "path tracing", "size": 16, "bold": True, "name": "Cambria", "color": INK, "after": 4},
            {"text": "o raio sai da câmera, atravessa um pixel da imagem e rebate na cena. em parede difusa o próximo rumo é aleatório; no espelho reflete; no vidro pode refratar. a cor do pixel é a luz acumulada nesse caminho.",
             "size": 13, "name": "Calibri", "color": INK, "after": 10},
            {"text": "monte carlo", "size": 16, "bold": True, "name": "Cambria", "color": INK, "after": 4},
            {"text": "estima um integral difícil (a luz que chega no pixel) pela média de amostras aleatórias. um caminho só é uma amostra ruim; milhares no mesmo pixel convergem para a imagem certa.",
             "size": 13, "name": "Calibri", "color": INK, "after": 10},
            {"text": "-w  -h  -s", "size": 16, "bold": True, "name": "Cambria", "color": INK, "after": 4},
            {"text": "-w = largura em pixels. -h = altura. -s = spp = amostras por subpixel (há 2×2 subpixels, então 4·spp caminhos por pixel). dobrar o spp dobra o tempo.",
             "size": 13, "name": "Calibri", "color": INK, "after": 0},
        ],
    )
    return slide


def make_renders(prs):
    slide = blank_slide(prs, 3, 3, "A APLICAÇÃO", "Isto é renderização: mais amostras, menos ruído")
    files = [
        (FIG / "render_spp1.png", "spp = 1", "4 caminhos/pixel. quase só ruído: a média ainda não estabilizou."),
        (FIG / "render_spp4.png", "spp = 4", "16 caminhos/pixel. a sala aparece, mas o granulado continua forte."),
        (FIG / "render_spp32.png", "spp = 32  (referência)", "128 caminhos/pixel. mesma entrada do relatório. T1 ≈ 54 s."),
    ]
    lefts = [600000, 4300000, 8000000]
    for x, (path, title, cap) in zip(lefts, files):
        add_pic(slide, path, x, 1950000, 3500000, 2800000)
        add_box(
            slide, x, 4800000, 3500000, 1600000,
            [
                {"text": title, "size": 15, "bold": True, "name": "Cambria", "color": INK, "after": 4},
                {"text": cap, "size": 13, "name": "Calibri", "color": INK, "after": 0},
            ],
        )
    return slide


def update_perfil(slide):
    wipe_body(slide, keep=4)
    for sh in slide.shapes:
        if sh.has_text_frame and "Onde o tempo" in sh.text_frame.text:
            fill_tf(
                sh.text_frame,
                [{"text": "Onde paralelizar — teto de Amdahl",
                  "size": 26, "bold": True, "name": "Cambria", "color": INK, "after": 0}],
            )
    add_pic(slide, FIG / "13_gprof.png", 500000, 1900000, 5800000, 4600000)
    add_rect(slide, 6500000, 1900000, 5200000, 4600000, WHITE, line=LINE)
    add_box(
        slide, 6700000, 2000000, 4850000, 4400000,
        [
            {"text": "o que o gráfico mostra", "size": 14, "bold": True, "name": "Cambria", "color": BLUE, "after": 4},
            {"text": "tempo próprio do gprof na versão sequencial. quase tudo (interseção, vetores, radiance, rng) é chamado de dentro do laço em y. só o main (~1,8%) fica de fora: setup e checksum.",
             "size": 12, "name": "Calibri", "color": INK, "after": 8},
            {"text": "o que é o teto de amdahl", "size": 14, "bold": True, "name": "Cambria", "color": BLUE, "after": 4},
            {"text": "se uma fração f é paralelizada e 1−f permanece serial, o speed-up com p threads não passa de S(p) = 1 / ((1−f) + f/p). a parte serial não encolhe, por mais threads que se coloque.",
             "size": 12, "name": "Calibri", "color": INK, "after": 8},
            {"text": "aqui f ≈ 0,982. teto em 6 threads ≈ 5,51; em 12 ≈ 10; com infinitas threads S∞ = 1/(1−f) ≈ 55,6. medimos 7,44× em 12 — abaixo de 10. a saturação não é essa fração serial.",
             "size": 12, "name": "Calibri", "color": INK, "after": 0},
        ],
    )


def update_estrategia(slide):
    wipe_body(slide, keep=4)
    add_pic(slide, FIG / "11_eixo_paralelismo.png",
            500000, 1850000, 11600000, 2800000)
    add_box(
        slide, 600000, 4700000, 5400000, 1900000,
        [
            {"text": "static (esquerda)", "size": 15, "bold": True, "name": "Cambria", "color": INK, "after": 4},
            {"text": "parte a imagem em faixas contíguas e fixas. a thread 0 fica com o chão (y=0), a última com o teto. se uma thread atrasa — no wsl2 isso acontece por preempção — o bloco inteiro dela trava e as outras não pegam o trabalho dela.",
             "size": 13, "name": "Calibri", "color": INK, "after": 0},
        ],
    )
    add_box(
        slide, 6200000, 4700000, 5400000, 1900000,
        [
            {"text": "dynamic,4 (direita)", "size": 15, "bold": True, "name": "Cambria", "color": INK, "after": 4},
            {"text": "as linhas entram num estoque comum em blocos de 4. a thread que termina pega o próximo bloco. o atraso de uma é absorvido pelas outras. foi a política adotada. cada linha continua independente: escrita só em c[y·w+x], rng privado, sem atomic.",
             "size": 13, "name": "Calibri", "color": INK, "after": 0},
        ],
    )


def update_diretivas(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("default(none)"):
            fill_tf(
                sh.text_frame,
                [
                    {"text": "default(none)", "size": 13, "bold": True, "name": "Calibri", "color": TEAL, "after": 2},
                    {"text": "obriga listar o escopo de cada variável. se algo deveria ser private e não foi declarado, o compilador recusa — não deixa a condição de corrida escondida.",
                     "size": 12, "name": "Calibri", "color": WHITE, "after": 8},
                    {"text": "shared", "size": 13, "bold": True, "name": "Calibri", "color": TEAL, "after": 2},
                    {"text": "só entradas (a, câmera), o buffer c (cada thread escreve numa faixa disjunta) e stat[tid] com padding de 64 B contra falso compartilhamento.",
                     "size": 12, "name": "Calibri", "color": WHITE, "after": 8},
                    {"text": "schedule(runtime) e nowait", "size": 13, "bold": True, "name": "Calibri", "color": TEAL, "after": 2},
                    {"text": "a política vem de OMP_SCHEDULE, então comparamos static/dynamic/guided sem recompilar. nowait tira a barreira do for: o tempo por thread é a carga real. sem reduction: não há acumulador no laço quente; o checksum é depois.",
                     "size": 12, "name": "Calibri", "color": WHITE, "after": 0},
                ],
            )


def update_politicas(slide):
    wipe_body(slide, keep=4)
    add_pic(slide, FIG / "16_schedule.png", 400000, 1900000, 6200000, 4600000)
    add_rect(slide, 6800000, 1900000, 4900000, 4600000, WHITE, line=LINE)
    add_box(
        slide, 7000000, 2000000, 4550000, 4400000,
        [
            {"text": "o que o gráfico mostra", "size": 14, "bold": True, "name": "Cambria", "color": BLUE, "after": 4},
            {"text": "mediana de 10 execuções em 12 threads. a haste é o intervalo min–max. barra mais baixa e haste curta = mais rápido e mais estável.",
             "size": 12, "name": "Calibri", "color": INK, "after": 8},
            {"text": "por que dynamic,4 (verde)", "size": 14, "bold": True, "name": "Cambria", "color": BLUE, "after": 4},
            {"text": "menor mediana (7,16 s) e menor dispersão (7,01–7,35 s). static sem chunk é o pior e o mais volátil (7,57–11,78 s): no wsl2 uma thread preemptada trava o bloco fixo. dynamic,4 redistribui e absorve o atraso.",
             "size": 12, "name": "Calibri", "color": INK, "after": 8},
            {"text": "o algoritmo em si desbalanceia < 1% entre threads. o fator decisivo não é o espelho vs. o fundo: é a robustez à preempção do sistema.",
             "size": 12, "name": "Calibri", "color": MUTED, "after": 0},
        ],
    )


def update_forte(slide):
    wipe_body(slide, keep=4)
    add_pic(slide, FIG / "14_speedup.png", 400000, 1880000, 6800000, 4600000)
    add_rect(slide, 7400000, 1880000, 4300000, 4600000, NAVY)
    add_box(
        slide, 7580000, 1980000, 4000000, 4400000,
        [
            {"text": "como ler o gráfico", "size": 14, "bold": True, "name": "Cambria", "color": TEAL, "after": 4},
            {"text": "eixo x = threads (1, 2, 4, 6, 8, 10, 12). linha cinza = speed-up ideal (y = p). azul = mediana medida. faixa clara = min/max das 10 execuções. linha pontilhada em 6 = início do smt.",
             "size": 12, "name": "Calibri", "color": WHITE, "after": 8},
            {"text": "até 6 núcleos físicos", "size": 14, "bold": True, "name": "Cambria", "color": TEAL, "after": 4},
            {"text": "speed-up 4,93 (teto de amdahl 5,51), eficiência 82,2%. a curva acompanha o ideal de perto.",
             "size": 12, "name": "Calibri", "color": WHITE, "after": 8},
            {"text": "depois de 6", "size": 14, "bold": True, "name": "Cambria", "color": TEAL, "after": 4},
            {"text": "duas threads no mesmo núcleo disputam ponto flutuante. em 12: 7,44× e 62%. 7,44 < teto 10 → a quebra é smt, não a fração serial.",
             "size": 12, "name": "Calibri", "color": PALE, "after": 0},
        ],
    )


def update_fraca(slide):
    wipe_body(slide, keep=4)
    add_pic(slide, FIG / "15_eficiencia.png", 400000, 1880000, 6800000, 4600000)
    add_rect(slide, 7400000, 1880000, 4300000, 4600000, WHITE, line=LINE)
    add_box(
        slide, 7580000, 1960000, 4000000, 4450000,
        [
            {"text": "como ler o gráfico", "size": 14, "bold": True, "name": "Cambria", "color": BLUE, "after": 4},
            {"text": "azul = eficiência da forte (trabalho fixo). amarelo = eficiência da fraca. 100% = o tempo não cresce ao colocar mais threads. a linha pontilhada em 6 marca o smt.",
             "size": 12, "name": "Calibri", "color": INK, "after": 8},
            {"text": "regra da fraca: spp = 32 · p", "size": 14, "bold": True, "name": "Cambria", "color": BLUE, "after": 4},
            {"text": "w e h ficam 800×600. o spp cresce com as threads porque o custo é linear em spp. o trabalho por thread permanece ~constante.",
             "size": 12, "name": "Calibri", "color": INK, "after": 8},
            {"text": "o que acontece", "size": 14, "bold": True, "name": "Cambria", "color": BLUE, "after": 4},
            {"text": "até 4 threads a fraca quase horizontal (55,6 → 58,6 s). em 12: 124 s (44,8%). cada núcleo recebe duas cargas completas, então a fraca cai mais que a forte. as duas quebram no mesmo ponto: a máquina, não o algoritmo.",
             "size": 12, "name": "Calibri", "color": INK, "after": 0},
        ],
    )


def update_conclusao(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("A saturação") or t.startswith("a saturação"):
            fill_tf(
                sh.text_frame,
                [
                    {"text": "o que os números dizem: de 54 s para 7,3 s (7,44×) numa máquina que já temos. o ponto doce são os 6 núcleos físicos (82,2%). as 12 threads lógicas ainda ajudam, mas com eficiência 62%.",
                     "size": 13, "name": "Calibri", "color": PALE, "after": 8},
                    {"text": "a saturação não é amdahl. f = 0,982 permitiria ~10× em 12 threads; medimos 7,44×. a causa é smt (duas threads disputam o ponto flutuante no mesmo núcleo) somada à preempção do wsl2. por isso adotamos dynamic,4: redistribui blocos quando uma thread atrasa.",
                     "size": 13, "name": "Calibri", "color": PALE, "after": 8},
                    {"text": "a versão paralela é equivalente à sequencial (checksum rgb idêntico de 1 a 12 threads e nas 7 políticas). tudo reproduzível: make all, scripts/run_all.sh e scripts/plot.py.",
                     "size": 13, "name": "Calibri", "color": PALE, "after": 0},
                ],
            )


def main():
    if not BACKUP.exists():
        shutil.copy2(PPTX, BACKUP)
    else:
        shutil.copy2(BACKUP, PPTX)

    ensure_render_pngs()
    prs = Presentation(str(PPTX))

    update_aplicacao(prs.slides[1])
    update_perfil(prs.slides[2])
    update_estrategia(prs.slides[3])
    update_diretivas(prs.slides[4])
    update_politicas(prs.slides[5])
    update_forte(prs.slides[6])
    update_fraca(prs.slides[7])
    update_conclusao(prs.slides[8])
    make_como_funciona(prs)
    make_renders(prs)

    for i, idx in enumerate(range(1, 11), start=1):
        set_badge_number(prs.slides[idx], i)

    prs.save(str(PPTX))
    print(f"salvo {PPTX}")
    print("slides", len(prs.slides))


if __name__ == "__main__":
    main()
